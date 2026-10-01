#!/usr/bin/env python3
"""
Probelauf — Entwicklungshilfe, nicht Teil der .exe
===================================================
Fährt den ganzen Weg gegen die Beispieldaten und prüft das Ergebnis gegen
festgeschriebene Sollwerte. Nach jeder Änderung einmal starten:

    python honorar_abrechner/probelauf.py

Die Sollwerte stammen aus dem Abgleich mit `Honorare_2025.xlsx` und den
beiden Musterbriefen. Weicht etwas ab, ist entweder ein Fehler entstanden
oder eine Regel hat sich bewusst geändert — dann gehören die Zahlen hier
angepasst, mit einem Satz dazu, warum.

Ohne die Beispieldaten (personenbezogen, nicht eingecheckt) überspringt
das Skript und sagt das.
"""

from __future__ import annotations

import shutil
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from honorar_abrechner import core          # noqa: E402

ALTMAPPE = Path(__file__).resolve().parent / "beispiele" / "Honorare_2025.xlsx"
JAHR = 2025

# --- Sollwerte, alle gegen die Altmappe belegt ------------------------
SOLL = {
    "empfaenger": 250,
    "buecher": 494,
    # Von den 68 Verträgen aus „Zahlung ab XX Ex.“ zahlen 56 ab einer festen
    # Stückzahl und laufen als Freimenge; nur die 12 ohne zahlenmäßige
    # Schwelle (zweite Auflage, Garantiehonorar, E-Book erst mit dem Buch)
    # bleiben „gesondert“. 2025 erreicht keiner der 56 seine Schwelle.
    "gesondert": 12,
    "stillgelegt": 105,
    # Sabine Maucher fällt heraus: ihre drei E-Books liegen unter zehn
    # Exemplaren, und die Zehnerregel aus der Notiz gilt jetzt.
    "briefe": 91,
    # 36,30 € weniger als zuvor: die Zehnerregel steht in 81 Notizen, war
    # aber nur bei 30 angekreuzt — der Verlag hat entschieden, dass sie
    # gilt (Maucher −15,36 €, Albertini −20,94 €). Dazu 7 Cent über der
    # Altmappe: der Satz je Exemplar wird jetzt immer auf
    # Cent gerundet. Schaupp „Freiheitsbäume“ 20 × 1,85 € = 37,00 € statt
    # 20 × 1,8458 € = 36,92 €; Scheidle einen Cent weniger.
    "auszahlung": 19105.95,
    "ksk": 14424.52,          # − 36,30 € Zehnerregel, siehe oben
    # Dieter Buck: steht so im Musterbrief und im alten KSK-Blatt.
    "buck_brutto": 2368.98,
    "buck_netto": 2214.01,
    "buck_posten": 16,
    # Drei Abweichungen bleiben, und sie sind EIN Sachverhalt: die
    # Summenkette AA245 der Altmappe fasst zwei Personen zusammen.
    "ohne_erklaerung": 3,
    # „D'accord mit de Welt“, zwei Titel mal zwei Mitautorinnen. Die
    # Altmappe rechnet =ROUND((17.9/1.07*0.05)/2,2), der Vertrag nennt
    # 10 % — 0,42 € statt 0,84 €. Bis der Verlag entscheidet, bleibt der
    # Betrag und die Meldung.
    "staffel_ohne_stand": 2,
}


class Probe:
    def __init__(self):
        self.fehler: list[str] = []

    def gleich(self, was: str, ist, soll, toleranz: float = 0.0):
        passt = (abs(ist - soll) <= toleranz if toleranz else ist == soll)
        print(f"   {'ok  ' if passt else 'FEHL'}  {was:<34} {ist!s:>12}"
              + ("" if passt else f"   erwartet: {soll}"))
        if not passt:
            self.fehler.append(f"{was}: {ist} statt {soll}")

    def wahr(self, was: str, bedingung: bool, erklaerung: str = ""):
        print(f"   {'ok  ' if bedingung else 'FEHL'}  {was}")
        if not bedingung:
            self.fehler.append(f"{was} — {erklaerung}" if erklaerung else was)


def main() -> int:
    if not ALTMAPPE.exists():
        print(f"Übersprungen: {ALTMAPPE.name} fehlt.")
        print("Die Beispieldaten sind personenbezogen und nicht eingecheckt.")
        return 0

    p = Probe()
    ordner = Path(tempfile.mkdtemp(prefix="honorar_probelauf_"))
    core.APP_DIR = ordner
    for name in ("CONFIG_PFAD", "BESTAND_PFAD", "SICHERUNG_PFAD",
                 "SPERR_PFAD", "WIEDERHERSTELLUNG_PFAD"):
        setattr(core, name, ordner / Path(getattr(core, name)).name)
    cfg = dict(core.DEFAULT_CONFIG)

    try:
        print("\n1. Import der Altmappe")
        bestand, protokoll = core.importiere_alt(ALTMAPPE, JAHR)
        p.gleich("Empfänger", len(bestand.empfaenger), SOLL["empfaenger"])
        p.gleich("Bücher", sum(1 for _ in bestand.buecher()), SOLL["buecher"])
        p.gleich("gesondert abzurechnen",
                 sum(1 for _, b in bestand.buecher() if b.gesondert),
                 SOLL["gesondert"])
        p.gleich("stillgelegt",
                 sum(1 for _, b in bestand.buecher() if b.stillgelegt),
                 SOLL["stillgelegt"])

        print("\n2. Speichern und wieder laden")
        core.speichere_bestand(bestand, jahr=JAHR, cfg=cfg)
        geladen = core.lade_bestand()
        p.wahr("keine Warnungen beim Laden", not geladen.warnungen,
               str(geladen.warnungen))
        vorher = core.rechne_alle(bestand, JAHR, cfg)
        nachher = core.rechne_alle(geladen, JAHR, cfg)
        p.gleich("Summe überlebt den Rundlauf",
                 round(sum(a.brutto for a in nachher if a.brief), 2),
                 round(sum(a.brutto for a in vorher if a.brief), 2), 0.005)

        print("\n3. Berechnung")
        alle = nachher
        p.gleich("Briefe", sum(1 for a in alle if a.brief), SOLL["briefe"])
        p.gleich("Auszahlungssumme",
                 round(sum(a.brutto for a in alle if a.brief), 2),
                 SOLL["auszahlung"], 0.005)
        p.gleich("an die KSK zu melden",
                 round(sum(a.ksk_netto for a in alle if a.ksk_netto > 0), 2),
                 SOLL["ksk"], 0.005)
        buck = next(a for a in alle if a.empfaenger.name == "Buck")
        p.gleich("Buck brutto", buck.brutto, SOLL["buck_brutto"], 0.005)
        p.gleich("Buck netto", buck.ksk_netto, SOLL["buck_netto"], 0.005)
        p.gleich("Buck Posten", len(buck.posten), SOLL["buck_posten"])
        p.wahr("kein Brief bei negativem Betrag",
               not any(a.brief and a.brutto < 0 for a in alle))
        # Vier Bücher rechnen mit der Hälfte der untersten Vertragsstufe,
        # weil die Altmappe „5 % und nochmal halbiert“ tippte. Der Betrag
        # bleibt, wie er war — aber die Meldung darf nicht verschwinden.
        gewarnt = {a.empfaenger.kennung for a in alle
                   for x in a.probleme if "Staffel vereinbart" in x}
        p.gleich("Warnung „Staffel ohne Stand“", len(gewarnt),
                 SOLL["staffel_ohne_stand"])

        # Staffelgrenze im Lauf des Jahres: 3450 vor dem Jahr, 73 dazu —
        # 50 zum alten, 23 zum neuen Satz. Kein Buch der Altmappe tut das
        # 2025, deshalb ein gebautes. Bei genau 3500 gehört das nächste
        # Exemplar schon zur höheren Stufe.
        test = core.Buch(kennung="B9999", titel="Staffeltest", kondition=(
            core.Kondition(ladenpreis=17.9, mwst_im_preis=7.0,
                           rabatt_anwenden=True, verlagsrabatt=40.0,
                           satz=12.0, staffel=[(3500, 12.0), (None, 13.0)])))
        test.jahr(JAHR - 1).vortrag = 3450
        test.jahr(JAHR).verkauft = 73
        posten = core.betrag_zeile(test, JAHR)
        p.gleich("Staffel im Jahr: Stufen",
                 [(m, pr) for m, _, pr, _ in posten.stufen],
                 [(50, 12.0), (23, 13.0)])
        p.gleich("Staffel im Jahr: Betrag", posten.netto, 89.90, 0.005)
        test.jahr(JAHR - 1).vortrag = 3500
        p.gleich("Staffel genau an der Grenze",
                 core.betrag_zeile(test, JAHR).satz_prozent, 13.0)

        print("\n4. Gegenprobe gegen die Altmappe")
        abweichungen = core.pruefe_gegen_excel(geladen, ALTMAPPE, JAHR, cfg)
        ohne = [a for a in abweichungen if not a.grund]
        p.gleich("Abweichungen ohne Erklärung", len(ohne),
                 SOLL["ohne_erklaerung"])
        for a in ohne:
            print(f"        {a.was}: {(a.wer or a.titel)[:42]} — "
                  f"alt {core.euro(a.soll)}, neu {core.euro(a.ist)}")

        print("\n5. Ausgaben")
        ziel = core.ausgabeordner(cfg, JAHR)
        briefe, ohne_brief = core.erzeuge_briefe(alle, cfg, ziel / "briefe")
        p.gleich("erzeugte Briefe", len(briefe), SOLL["briefe"])
        from docx import Document
        spalten = {len(Document(str(b)).tables[-1].columns) for b in briefe}
        leer = [b.name for b in briefe
                if len(Document(str(b)).tables[-1].rows) <= 1]
        p.wahr("kein Brief mit leerer Titeltabelle", not leer, str(leer[:3]))
        p.wahr("MwSt-Spalte wird entfernt, wo sie nicht gebraucht wird",
               spalten == {7, 8}, f"Spaltenzahlen: {sorted(spalten)}")
        listen = core.schreibe_listen(alle, JAHR,
                                      ziel / f"Zahlungsliste_{JAHR}.xlsx", cfg)
        import openpyxl
        blaetter = openpyxl.load_workbook(listen).sheetnames
        p.wahr("drei Listen in einer Datei", len(blaetter) == 3, str(blaetter))

        print("\n6. Zahlen lesen")
        for text, soll in [("1.234", 1234), ("1 234", 1234), ("12,5", None),
                           ("abc", None), ("-9", -9), ("0", 0), ("", None)]:
            zahl, _ = core.lies_stueckzahl(text)
            p.gleich(f"lies_stueckzahl({text!r})", zahl, soll)
        p.gleich("kaufmännisch runden: 5,635", core.runde(5.635), 5.64, 0.0001)
        p.gleich("kaufmännisch runden: 2,675", core.runde(2.675), 2.68, 0.0001)

        print("\n" + "=" * 58)
        if p.fehler:
            print(f"   {len(p.fehler)} Abweichung(en) vom Sollwert:")
            for f in p.fehler:
                print(f"      · {f}")
            return 1
        print("   Alles wie erwartet.")
        return 0
    finally:
        shutil.rmtree(ordner, ignore_errors=True)


if __name__ == "__main__":
    sys.exit(main())
